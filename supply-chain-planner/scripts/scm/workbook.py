"""Public-dependency XLSX export. Core reports remain the calculation authority."""
from datetime import date
from io import BytesIO
from pathlib import Path

from plan_inventory import ContractError
from .contracts import load_json, num, provenance, text


def export(plan_path, output, backtest_path=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.workbook.properties import CalcProperties
    except ImportError as e:
        raise ContractError("portable XLSX export requires public package openpyxl>=3.1") from e
    report = load_json(plan_path)
    if not isinstance(report, dict) or not isinstance(report.get("items"), list) or not report["items"]:
        raise ContractError("export: require an inventory plan report.json with nonempty items")
    target = Path(output)
    if target.suffix.lower() != ".xlsx":
        raise ContractError("export: output must have .xlsx extension")
    if target.exists():
        raise ContractError("export: output already exists")
    book = Workbook()
    book.remove(book.active)
    book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)

    def sheet(name, headers, rows_, note):
        ws = book.create_sheet(name)
        width = len(headers)
        ws.append([name])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
        ws.append([note])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        ws.append(headers)
        for r in rows_:
            ws.append(r)
        for row in ws:
            for c in row:
                if isinstance(c.value, str):
                    c.data_type = "s"  # User strings are never formulas, including = + - @ prefixes.
                c.font = Font(name="Calibri", size=11, color="243145")
                c.alignment = Alignment(vertical="center", wrap_text=True)
                if isinstance(c.value, (int, float)):
                    c.number_format = "#,##0.00"
                if isinstance(c.value, date):
                    c.number_format = "yyyy-mm-dd"
        for c in ws[1]:
            c.font = Font(name="Calibri", size=18, bold=True, color="203B61")
        for c in ws[3]:
            c.fill = PatternFill("solid", fgColor="203B61")
            c.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
        for i in range(4, ws.max_row + 1):
            ws.row_dimensions[i].height = 32
            for c in ws[i]:
                if i % 2 == 0:
                    c.fill = PatternFill("solid", fgColor="F2F5F9")
        ws.row_dimensions[1].height = 32
        ws.row_dimensions[2].height = 38
        ws.row_dimensions[3].height = 32
        for col in ws.columns:
            letter = col[2].column_letter
            ws.column_dimensions[letter].width = 20 if col[2].column > 2 else 24
        ws.freeze_panes = "C4"
        ws.auto_filter.ref = f"A3:{ws.cell(ws.max_row, width).coordinate}"
        ws.sheet_view.showGridLines = False
        return ws

    main_rows, projections = [], []
    for item in report["items"]:
        sku, loc = text(item.get("sku"), "sku"), text(item.get("location"), "location")
        q, cost = item.get("suggested_order_qty"), item.get("unit_cost")
        if q is not None:
            num(q, "suggested_order_qty", whole=True)
        if cost is not None:
            num(cost, "unit_cost")
        main_rows.append([sku, loc, item.get("recommendation_state"), q, cost,
                          item.get("currency"), None, item.get("history_missing_days")])
        for row in item.get("daily_projection", []):
            r = row["with_order"]
            projections.append([sku, loc, date.fromisoformat(row["date"]), r["demand"], r["received"],
                                r["ending_stock"], r["ending_backlog"], r["lost_units"]])
    source = provenance(plan_path)
    note = f"快照 {report.get('parameters', {}).get('as_of', '未提供')}；引擎 {report.get('engine_version', '未提供')}；{source['input_name']} SHA256 {source['input_sha256']}"
    ws = sheet("补货计划", ["SKU", "仓库", "复核状态", "建议数量/件", "单位成本", "币种", "采购货值", "历史缺日数"], main_rows, note)
    for i in range(4, ws.max_row + 1):
        ws.cell(i, 7, f'=IF(AND(ISNUMBER(D{i}),ISNUMBER(E{i})),D{i}*E{i},"")')
        for j in (1, 2):
            ws.cell(i, j).number_format = "@"
        ws.cell(i, 4).number_format = "#,##0"
        ws.cell(i, 8).number_format = "#,##0"
    sheet("逐日预测", ["SKU", "仓库", "日期", "需求/件", "到货/件", "日末实物/件", "日末欠单/件", "失销/件"],
          projections, "条件点预测；到货先于当日需求。修改源数据后重新运行引擎；表内只重算采购货值。")
    if backtest_path:
        bt = load_json(backtest_path)
        if not isinstance(bt, dict) or bt.get("kind") != "policy_backtest":
            raise ContractError("export: invalid policy_backtest report")
        runs = bt["test_runs"]
        rows_ = [[r["policy"], r["seed"], r["immediate_fill_rate"], r["modeled_operating_cost"],
                  r["purchase_cash"], r["ending_stock"], r["ending_backlog"], r["ending_pipeline_units"]] for r in runs]
        sb = sheet("政策回测", ["政策", "交期种子", "即时满足率", "模型运营成本", "采购现金", "期末库存", "期末欠单", "期末在途"], rows_,
                   f"测试期 {bt['split']['test'][0]} 至 {bt['split']['test'][1]}；采用建议 {bt['adoption_gate']['recommendation']}；{provenance(backtest_path)['input_name']}")
        for i in range(4, sb.max_row + 1):
            sb.cell(i, 3).number_format = "0.0%"
    target.parent.mkdir(parents=True, exist_ok=True)
    buffer = BytesIO()
    book.save(buffer)
    with target.open("xb") as stream:
        stream.write(buffer.getvalue())
    return {"status": "created", "sheets": book.sheetnames, "rows": len(main_rows),
            "formula_calculation": "Excel recalculates on open; exporter does not evaluate formulas"}

"""CSV/Excel adapters, explicit field mapping, and validated scenario settings."""
import csv
import hashlib
import json
import math
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path

from plan_inventory import (Config, ContractError, ITEM_FIELDS, DEMAND_FIELDS, INBOUND_FIELDS,
                            load_inputs, number, parse_date, read_table)

ALIASES = {"货号": "sku", "SKU": "sku", "仓库": "location", "在手库存": "on_hand",
           "已占用": "reserved", "欠单": "backorders", "交期天数": "lead_time_days",
           "复核周期": "review_period_days", "服务目标": "service_level", "起订量": "moq",
           "箱规": "pack_size", "单位成本": "unit_cost", "币种": "currency",
           "日期": "date", "需求量": "demand_qty", "缺货标记": "stockout",
           "订单行号": "order_id", "未交数量": "qty", "预计到货日": "eta", "状态": "status",
           "单位日持有成本": "holding_cost_per_unit_day", "缺货成本": "shortage_cost",
           "单次订货成本": "order_fixed_cost", "上市日期": "launch_date"}
SHEETS = {"items": ("items", "商品", "库存"), "demand": ("demand", "需求"),
          "inbound": ("inbound", "在途")}
ECONOMICS = ("holding_cost_per_unit_day", "shortage_cost", "order_fixed_cost")
DEFAULTS = {"history_days": 180, "demand_kind": "orders", "missing_days": "error",
            "shortage_mode": "backorder", "demand_factor": 1.,
            "lead_time_scenarios": [{"delay_days": 0, "probability": 1.}],
            "budget": None, "capacity_units": None, "currency": None,
            "allocation": "service", "forecast_method": None, "safety_scale": 1.}


def settings(path):
    raw = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(raw, dict) or set(raw) - (set(DEFAULTS) | {"as_of"}):
        raise ContractError("settings: expected object with documented keys only")
    c = {**DEFAULTS, **raw}
    c["as_of"] = parse_date(c.get("as_of"), "settings:as_of")
    c["history_days"] = number(c["history_days"], "history_days", integer=True, low=7, high=730)
    c["demand_factor"] = number(c["demand_factor"], "demand_factor", low=.01, high=100)
    c["safety_scale"] = number(c["safety_scale"], "safety_scale", low=0, high=3)
    for key, allowed in (("demand_kind", ("orders", "sales")), ("missing_days", ("error", "zero")),
                         ("shortage_mode", ("backorder", "lost_sales")),
                         ("allocation", ("service", "economic"))):
        if c[key] not in allowed:
            raise ContractError(f"settings:{key}: expected {allowed}")
    from .forecast import METHODS
    if c["forecast_method"] is not None and c["forecast_method"] not in METHODS:
        raise ContractError("settings: unknown forecast_method")
    for key in ("budget", "capacity_units"):
        if c[key] is not None:
            c[key] = number(c[key], key)
    scenarios = c["lead_time_scenarios"]
    if not isinstance(scenarios, list) or not 1 <= len(scenarios) <= 10:
        raise ContractError("lead_time_scenarios: need 1..10 explicit delay/probability pairs")
    checked = []
    for s in scenarios:
        if not isinstance(s, dict) or set(s) != {"delay_days", "probability"}:
            raise ContractError("lead scenario requires exactly delay_days and probability")
        checked.append((number(s["delay_days"], "delay_days", integer=True, high=90),
                        number(s["probability"], "probability", low=1e-12, high=1)))
    if len({x[0] for x in checked}) != len(checked) or not math.isclose(sum(x[1] for x in checked), 1., abs_tol=1e-9):
        raise ContractError("lead delays must be unique and probabilities must sum to 1")
    c["lead_scenarios"] = sorted(checked)
    if c["allocation"] == "service" and (c["budget"] is not None or c["capacity_units"] is not None):
        raise ContractError("shared constraints require allocation=economic and explicit cost inputs")
    return c


def canonical_headers(headers, table, mapping=None):
    if any(h is None or str(h).strip() == "" for h in headers):
        raise ContractError(f"{table}: blank header")
    explicit = (mapping or {}).get("columns", {}).get(table, {})
    mapped = [explicit.get(str(h).strip(), ALIASES.get(str(h).strip(), str(h).strip())) for h in headers]
    if len(mapped) != len(set(mapped)):
        raise ContractError(f"{table}: multiple source columns map to the same field")
    return mapped


def excel_to_csv(path, out, mapping=None):
    """Read only. Formula cells are rejected instead of trusting stale caches."""
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError("Excel import requires openpyxl") from e
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except (openpyxl.utils.exceptions.InvalidFileException, zipfile.BadZipFile, ValueError) as e:
        raise ContractError(f"Excel: invalid or unsupported workbook: {e}") from e
    audit, paths = [], {}
    try:
        for table, aliases in SHEETS.items():
            chosen = (mapping or {}).get("sheets", {}).get(table)
            if chosen is None:
                matches = [s for s in book.sheetnames if s in aliases]
                if len(matches) != 1:
                    raise ContractError(f"Excel:{table}: expected exactly one matching sheet or explicit mapping")
                chosen = matches[0]
            if chosen not in book.sheetnames:
                raise ContractError(f"Excel: missing sheet {chosen}")
            rows = iter(book[chosen].iter_rows())
            first = next(rows, None)
            if not first:
                raise ContractError(f"Excel:{chosen}: missing header")
            if any(c.data_type == "f" for c in first):
                raise ContractError(f"Excel:{chosen}: formula header")
            headers = canonical_headers([c.value for c in first], table, mapping)
            target = out / f"{table}.csv"
            with target.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(headers)
                for row in rows:
                    if all(c.value is None for c in row):
                        continue
                    values = []
                    if any(c.value is not None for c in row[len(headers):]):
                        raise ContractError(f"Excel:{chosen}: data extends beyond the declared headers")
                    for index, field in enumerate(headers):
                        if index >= len(row):
                            values.append("")
                            continue
                        cell = row[index]
                        v = cell.value
                        if cell.data_type in ("f", "e"):
                            raise ContractError(f"Excel:{chosen}:{cell.coordinate}: formula/error cell; provide verified values")
                        if field in ("sku", "location", "order_id") and v is not None and not isinstance(v, str):
                            raise ContractError(f"Excel:{chosen}:{cell.coordinate}: identifiers must be TEXT to preserve leading zeros")
                        if isinstance(v, datetime):
                            if v.time().isoformat() != "00:00:00":
                                raise ContractError(f"Excel:{chosen}:{cell.coordinate}: dates must not contain times")
                            v = v.date().isoformat()
                        elif isinstance(v, date):
                            v = v.isoformat()
                        elif isinstance(v, bool):
                            raise ContractError(f"Excel:{chosen}:{cell.coordinate}: boolean is not a quantity")
                        values.append("" if v is None else v)
                    writer.writerow(values)
            audit.extend({"table": table, "sheet": chosen, "source_column": str(a.value),
                          "canonical_field": b} for a, b in zip(first, headers))
            paths[table] = target
    finally:
        book.close()
    return paths, audit


def load_bundle(c, *, workbook=None, items=None, demand=None, inbound=None, mapping=None):
    if mapping is not None:
        if not isinstance(mapping, dict) or set(mapping) - {"sheets", "columns"}:
            raise ContractError("mapping: expected sheets/columns object")
        for field in ("sheets", "columns"):
            if not isinstance(mapping.get(field, {}), dict) or set(mapping.get(field, {})) - set(SHEETS):
                raise ContractError(f"mapping:{field}: expected documented table keys")
        if any(not isinstance(k, str) or not isinstance(v, str) for k, v in mapping.get("sheets", {}).items()):
            raise ContractError("mapping:sheets: sheet names must be strings")
        for values in mapping.get("columns", {}).values():
            if not isinstance(values, dict) or any(not isinstance(k, str) or not isinstance(v, str) or not v for k, v in values.items()):
                raise ContractError("mapping:columns: expected nonempty source/target field names")
    sources = {"workbook": workbook} if workbook else {"items": items, "demand": demand, "inbound": inbound}
    if not all(sources.values()):
        raise ContractError("provide a workbook or all three CSV files, including empty inbound header")
    hashes = {k: {"name": Path(v).name, "sha256": hashlib.sha256(Path(v).read_bytes()).hexdigest()}
              for k, v in sources.items()}
    with tempfile.TemporaryDirectory(prefix="scm-import-") as temp:
        if workbook:
            paths, audit = excel_to_csv(workbook, Path(temp), mapping)
        else:
            paths, audit = sources, []
        config = Config(c["as_of"], c["history_days"], c["demand_kind"], c["missing_days"])
        products, history, receipts = load_inputs(paths["items"], paths["demand"], paths["inbound"], config)
        economics = {}
        for _, row in read_table(paths["items"], ITEM_FIELDS):
            vals = [row.get(field, "") for field in ECONOMICS]
            if any(vals) and not all(vals):
                raise ContractError("all three operating-cost fields must be supplied together")
            if all(vals):
                if not row.get("currency"):
                    raise ContractError("operating-cost inputs need currency and unit_cost")
                economics[(row["sku"], row["location"])] = {
                    field: number(value, field) for field, value in zip(ECONOMICS, vals)}
        if any(p.lead_time_days + p.review_period_days + max(d for d, _ in c["lead_scenarios"]) > 366 for p in products):
            raise ContractError("lead + maximum delay + review must be <=366")
        if c["shortage_mode"] == "lost_sales" and any(p.backorders for p in products):
            raise ContractError("lost_sales cannot silently discard historical backorders")
        if c["allocation"] == "economic":
            if any(p.key not in economics or p.unit_cost is None for p in products):
                raise ContractError("economic allocation needs unit_cost, currency and all operating costs for every item")
            currencies = {p.currency for p in products}
            if len(currencies) != 1 or c["currency"] not in currencies:
                raise ContractError("economic allocation requires one explicit, matching currency")
    return products, history, receipts, economics, {"inputs": hashes, "field_mapping": audit}


def json_safe(value):
    if isinstance(value, dict):
        return {k: json_safe(v) for k, v in value.items()}
    if isinstance(value, (tuple, list)):
        return [json_safe(x) for x in value]
    if isinstance(value, date):
        return value.isoformat()
    return value

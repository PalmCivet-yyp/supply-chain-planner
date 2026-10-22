#!/usr/bin/env python3
"""Fixed synthetic acceptance checks. Unsupported workflows score separately from wrong answers."""
import argparse
import copy
import hashlib
import importlib
import json
import math
import sys
import tempfile
from dataclasses import replace
from datetime import date, timedelta
from pathlib import Path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skill-dir", required=True)
    parser.add_argument("--out", required=True, help="new result JSON")
    args = parser.parse_args()
    skill = Path(args.skill_dir).resolve()
    sys.path.insert(0, str(skill / "scripts"))
    from plan_inventory import Item, Receipt, ContractError
    from scm.engine import plan
    from scm.backtest import backtest
    from scm import VERSION
    fixtures_path = Path(__file__).resolve().parents[1] / "assets/evaluation/fixtures.json"
    fixture = json.loads(fixtures_path.read_text())
    cases = fixture["cases"]
    day = date(2026, 10, 20)
    econ = {"holding_cost_per_unit_day": .1, "shortage_cost": 3., "order_fixed_cost": 2.}
    config = {"as_of": day, "history_days": 180, "missing_days": "error", "demand_kind": "orders",
              "shortage_mode": "backorder", "lead_scenarios": [(0, 1.)], "demand_factor": 1.,
              "allocation": "service", "forecast_method": None, "budget": None, "capacity_units": None,
              "currency": "CNY", "safety_scale": 1.}
    def item(**changes):
        launch = changes.pop("launch_date", None)
        p = replace(Item("001", "GZ", 10, 0, 0, 2, 3, .95, 6, 3, 2., "CNY"), **changes)
        # The baseline CSV loader silently ignored this extra column; mirror that behavior.
        return replace(p, launch_date=launch) if "launch_date" in p.__dataclass_fields__ else p
    def history(days=180, value=10):
        return {day - timedelta(days=i): (value, 0) for i in range(days)}
    def eq(a, b):
        if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) and not isinstance(b, bool):
            assert math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-9), (a, b)
        else:
            assert a == b, (a, b)
    def reject(fn):
        try:
            fn()
        except (ContractError, ValueError):
            return
        raise AssertionError("invalid input was accepted")
    def inventory(case):
        p, h, c = item(), history(), dict(config)
        if case in ("C03", "C04", "C05", "C06", "C16"):
            h = history(60, 4)
        if case == "C03":
            h[day - timedelta(days=200)] = (4, 0)
        if case in ("C05", "C16"):
            p = item(launch_date=day - timedelta(days=59 if case == "C05" else 58))
        if case == "C06":
            p = item(launch_date=day - timedelta(days=60))
        if case == "C02":
            del h[day - timedelta(days=10)]
        if case == "C08":
            c["demand_kind"] = "sales"; h[day] = (1, 1)
        if case == "C09":
            h = history(value=0); p = item(on_hand=0)
        receipts = [Receipt("P", *p.key, 500, day + timedelta(days=40), "confirmed")] if case == "C07" else []
        if case == "C10":
            p = item(on_hand=0, lead_time_days=3)
        if case == "C11":
            q = item(location="SH", on_hand=100)
            r = plan([p, q], {p.key: h, q.key: history(value=1)}, [], {}, c)
            values = {(x["sku"], x["location"]): x["suggested_order_qty"] for x in r["items"]}
            eq(values[p.key], 42); eq(values[q.key], 0); return
        a = plan([p], {p.key: h}, receipts, {}, c)
        r = a["items"][0]
        if case == "C01": eq(r["suggested_order_qty"], 42)
        if case in ("C02", "C03", "C06", "C08", "C16"): eq(r["recommendation_state"], "blocked")
        if case == "C04": eq(r["recommendation_state"], "scenario_only")
        if case == "C05": eq(r["recommendation_state"], "ready_for_review")
        if case in ("C07", "C09"): eq(r["suggested_order_qty"], 0)
        if case == "C10":
            b = plan([p], {p.key: h}, [], {}, {**c, "shortage_mode": "lost_sales"})["items"][0]
            eq(r["daily_projection"][2]["with_order"]["ending_stock"], 30)
            eq(b["daily_projection"][2]["with_order"]["ending_stock"], 50)
        if case == "C12":
            h[day + timedelta(days=1)] = (999999, 0)
            eq(a, plan([p], {p.key: h}, [], {}, c))
    bt_cache = {}
    def replay(case):
        p, h = item(review_period_days=7), history(300, 4)
        if "normal" not in bt_cache:
            bt_cache["normal"] = backtest([p], {p.key: h}, {p.key: econ}, config, 120, 56, 56, (17,))[0]
        a = bt_cache["normal"]
        if case == "C13":
            access = a.get("history_access", {})
            eq(access.get("first_review_requested_start"), "2026-01-02")
            eq(access.get("first_review_actual_starts", [{}])[0].get("start"), "2026-01-02")
            eq(access.get("warmup_before_training_allowed"), True)
        if case == "C14":
            h[day] = (9999, 0)
            b = backtest([p], {p.key: h}, {p.key: econ}, config, 120, 56, 56, (17,))[0]
            eq(a["validation"], b["validation"]); eq(a["selected_adaptive"], b["selected_adaptive"])
        if case == "C15":
            ds = a["decision_audits"][0]["decisions"]
            assert ds and all(d.get("history_windows") for d in ds), "missing per-decision history dates"
            assert all(w["start"] <= w["end"] <= d["as_of"] for d in ds for w in d["history_windows"])
    def operational(case, category):
        names = {"procurement": ("procurement", "analyze"), "sourcing": ("sourcing", "compare"), "sop": ("sop", "analyze")}
        module, fn = names[category]
        if not (skill / f"scripts/scm/{module}.py").exists():
            return "unsupported"
        analyze = getattr(importlib.import_module("scm." + module), fn)
        data = copy.deepcopy(fixture["inputs"][category])
        if category == "procurement":
            if case == "P02":
                data["spend"][1]["supplier"] = "A"
            if case == "P05":
                data["spend"][0].pop("upstream_group")
            if case == "P08":
                data["deliveries"].append({**data["deliveries"][0], "id": "future", "due_date": "2026-11-10"})
            if case == "P10": data["matches"][0]["invoice_amount"] = 126
            if case == "P11": data["spend"].append(data["spend"][0].copy())
            if case == "P12": data["cycles"][0]["approved"] = "2026-08-01"
            if case in ("P11", "P12"): reject(lambda: analyze(data)); return
            r = analyze(data)
            g = next(x for x in r["spend_groups"] if x["currency"] == "CNY")
            if case == "P01":
                eq(g["total_spend"], 1000); eq(len(r["spend_groups"]), 2)
            if case == "P02": eq(g["supplier_count"], 1)
            if case == "P03": eq(g["hhi_fraction"], .6**2 + .4**2)
            if case == "P04":
                eq(g["upstream_group_count"], 1); eq(g["independence_status"], "shared_upstream_exposure")
            if case == "P05": eq(g["upstream_group_count"], None); eq(g["independence_status"], "unknown")
            if case == "P06":
                m = next(x for x in r["cycle_metrics"] if x["stage"] == "request_to_po")
                eq(m["median_days"], 4); eq(m["p90_days"], 4)
            if case == "P07": eq(r["supplier_otif"][0]["line_otif"], 0)
            if case == "P08": eq(r["supplier_otif"][0]["due_line_count"], 2)
            if case == "P09": eq(r["three_way_matches"][0]["status"], "partial_match"); eq(r["three_way_matches"][0]["expected_invoice_subtotal"], 125)
            if case == "P10": eq(r["three_way_matches"][0]["status"], "discrepancy")
        if category == "sourcing":
            if case in ("T03", "T04"): data["options"][1]["requirements"]["warranty"] = "fail" if case == "T03" else "unknown"
            if case == "T05":
                for s in data["sources"]: s["status"] = "unverified"
            if case == "T06":
                for s in data["sources"]: s["checked_on"] = "2025-01-01"
            if case == "T07": data["sources"][1]["status"] = "refuted"
            if case == "T08": data["sources"][0]["checked_on"] = "2026-10-21"
            if case == "T09": data["options"][1]["currency"] = "USD"
            if case == "T10": del data["options"][1]["costs"]["annual_support"]
            if case == "T11": data["criteria"][0]["weight"] = .7
            if case == "T12":
                data["criteria"][0]["weight"] = data["criteria"][1]["weight"] = .5
                data["options"][0]["scores"]["quality"] = 100
                data["options"][1]["scores"]["quality"] = 0
            if case in ("T08", "T09", "T10", "T11"): reject(lambda: analyze(data)); return
            r = analyze(data)
            if case == "T01": eq(r["options"][0]["tco"], 1620); eq(r["options"][1]["tco"], 1320)
            if case == "T02": eq(r["leading_option_for_review"], "B")
            if case in ("T03", "T04", "T07"): eq(r["leading_option_for_review"], "A")
            if case in ("T05", "T06"): eq(r["leading_option_for_review"], None)
            if case == "T12": eq({s["leading_option"] for s in r["sensitivity"]}, {"A", "B"})
        if category == "sop":
            if case == "O02":
                for h in data["history"]: h["actual"] = h["forecast"] = 0
            if case == "O03": data["history"][0]["actual"] = 0
            if case == "O07":
                data["opening"].append({**data["opening"][0], "location": "SH", "stock": 1000})
                data["plan"] += [{**x, "location": "SH"} for x in list(data["plan"])]
            if case == "O08": data["owners"].pop("supply")
            if case == "O09": data["history"][0]["period"] = "2026-11"
            if case == "O10": data["plan"].append(data["plan"][0].copy())
            if case == "O11": data["plan"][1]["period"] = "2027-01"
            if case == "O12":
                for p in data["plan"]: p["planned_supply"] = 0
            if case in ("O09", "O10", "O11"): reject(lambda: analyze(data)); return
            r = analyze(data)
            base = [x for x in r["balance"] if x["scenario"] == "base" and x["location"] == "GZ"]
            if case == "O01": eq(r["forecast_metrics"]["wape"], .1); eq(r["forecast_metrics"]["bias_fraction"], -10 / 300)
            if case == "O02": eq(r["forecast_metrics"]["wape"], None); eq(r["forecast_metrics"]["mape_nonzero_actual"], None)
            if case == "O03": eq(r["forecast_metrics"]["mape_coverage"], .5); eq(r["forecast_metrics"]["mape_nonzero_actual"], .1)
            if case == "O04": eq([x["ending_backlog"] for x in base], [30, 20])
            if case == "O05": eq(base[1]["effective_supply"], 90); eq(base[1]["capacity_excess"], 10)
            if case == "O06": eq([x["ending_backlog"] for x in r["balance"] if x["scenario"] == "demand_up_20pct"], [50, 56])
            if case == "O07": eq([x["ending_backlog"] for x in r["balance"] if x["scenario"] == "base" and x["location"] == "SH"], [0, 0])
            if case == "O08": assert r["actions"] and all(x["status"] == "awaiting_owner" for x in r["actions"])
            if case == "O12": eq([x["ending_backlog"] for x in base], [80, 160])
    def workbook(case):
        if not (skill / "scripts/scm/workbook.py").exists(): return "unsupported"
        from scm.workbook import export
        import openpyxl
        p = item(sku="=2+2" if case == "X03" else "001")
        r = plan([p], {p.key: history()}, [], {}, config)
        if case == "X04": r["items"][0]["suggested_order_qty"] = None
        if case == "X05": r["items"][0]["unit_cost"] = None
        if case == "X06": r["items"][0]["daily_projection"] = []
        with tempfile.TemporaryDirectory() as t:
            root = Path(t); ip = root / "plan.json"; op = root / "output.xlsx"
            ip.write_text(json.dumps(r))
            bp = None
            if case == "X08":
                bp = root / "backtest.json"
                bp.write_text(json.dumps(bt_cache.get("normal") or backtest([item(review_period_days=7)],
                    {item().key: history(300, 4)}, {item().key: econ}, config, 120, 56, 56, (17,))[0]))
            export(ip, op, bp)
            if case == "X07":
                original = op.read_bytes(); reject(lambda: export(ip, op)); eq(op.read_bytes(), original); return
            b = openpyxl.load_workbook(op, data_only=False)
            try:
                ws = b["补货计划"]
                if case == "X01":
                    eq(b.sheetnames[:2], ["补货计划", "逐日预测"])
                    eq(ws["G4"].data_type, "f")
                    assert "D4" in ws["G4"].value and "E4" in ws["G4"].value
                if case == "X02": eq(ws["A4"].value, "001"); eq(ws["A4"].data_type, "s")
                if case == "X03": eq(ws["A4"].value, "=2+2"); eq(ws["A4"].data_type, "s")
                if case == "X04": eq(ws["D4"].value, None)
                if case == "X05": eq(ws["E4"].value, None)
                if case == "X06": eq(b["逐日预测"].max_row, 3)
                if case == "X08": assert isinstance(b["政策回测"]["C4"].value, (int, float))
            finally: b.close()
    output = []
    for case in cases:
        name, category = case["id"], case["category"]
        try:
            status = (replay(name) if name in ("C13", "C14", "C15") else inventory(name)) if category == "inventory" else workbook(name) if category == "export" else operational(name, category)
            output.append({**case, "status": status or "pass"})
        except Exception as e:
            output.append({**case, "status": "fail", "detail": f"{type(e).__name__}: {e}"})
    summary = {c: {s: sum(x["category"] == c and x["status"] == s for x in output) for s in ("pass", "fail", "unsupported")} for c in dict.fromkeys(x["category"] for x in cases)}
    result = {"kind": "synthetic_acceptance", "target_version": VERSION,
              "fixture_sha256": hashlib.sha256(fixtures_path.read_bytes()).hexdigest(),
              "runner_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
              "scope": "Developer acceptance cases; no real-business savings or independent forecast benchmark.",
              "summary": summary, "cases": output}
    with Path(args.out).open("x", encoding="utf-8") as stream:
        json.dump(result, stream, ensure_ascii=False, indent=2)
    print(json.dumps({"version": VERSION, "summary": summary}))
    return 1 if any(x["status"] == "fail" for x in output) else 0


if __name__ == "__main__":
    raise SystemExit(main())
